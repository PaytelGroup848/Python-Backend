from pathlib import Path

from openpyxl import load_workbook

from app.modules.ingestion.parsers.base_parser import (
    BaseParser
)

from app.modules.ingestion.schemas.parsed_document import (
    ParsedDocument
)


class XlsxParser(
    BaseParser
):

    @property
    def parser_code(
        self
    ) -> str:

        return "XLSX"

    @property
    def supported_extensions(
        self
    ) -> set[str]:

        return {
            ".xlsx"
        }

    @property
    def supported_mime_types(
        self
    ) -> set[str]:

        return {
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        }

    async def parse(
        self,
        file_path: Path,
        configuration: dict | None = None
    ) -> ParsedDocument:

        workbook = load_workbook(

            filename=file_path,

            data_only=True

        )

        extracted_rows = []

        sheet_names = []

        total_rows = 0

        for worksheet in workbook.worksheets:

            sheet_names.append(
                worksheet.title
            )

            extracted_rows.append(
                f"# Sheet: {worksheet.title}"
            )

            for row in worksheet.iter_rows(
                values_only=True
            ):

                values = [

                    str(value)

                    if value is not None

                    else ""

                    for value in row

                ]

                extracted_rows.append(
                    " | ".join(values)
                )

                total_rows += 1

            extracted_rows.append("")

        text_content = "\n".join(
            extracted_rows
        )

        return ParsedDocument(

            file_name=file_path.name,

            file_extension=file_path.suffix.lower(),

            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",

            parser_code=self.parser_code,

            title=file_path.stem,

            language=None,

            page_count=None,

            character_count=len(
                text_content
            ),

            text_content=text_content,

            metadata={

                "sheet_count": len(
                    workbook.sheetnames
                ),

                "sheet_names": sheet_names,

                "row_count": total_rows

            }

        )


xlsx_parser = XlsxParser()