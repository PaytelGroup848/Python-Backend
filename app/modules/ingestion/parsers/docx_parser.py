from pathlib import Path

from docx import Document

from app.modules.ingestion.parsers.base_parser import (
    BaseParser
)

from app.modules.ingestion.schemas.parsed_document import (
    ParsedDocument
)


class DocxParser(
    BaseParser
):

    @property
    def parser_code(
        self
    ) -> str:

        return "DOCX"

    @property
    def supported_extensions(
        self
    ) -> set[str]:

        return {
            ".docx"
        }

    @property
    def supported_mime_types(
        self
    ) -> set[str]:

        return {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        }

    async def parse(
        self,
        file_path: Path,
        configuration: dict | None = None
    ) -> ParsedDocument:

        document = Document(
            file_path
        )

        paragraphs = []

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:

                paragraphs.append(
                    text
                )

        text_content = "\n".join(
            paragraphs
        )

        metadata = document.core_properties

        return ParsedDocument(

            file_name=file_path.name,

            file_extension=file_path.suffix.lower(),

            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",

            parser_code=self.parser_code,

            title=metadata.title,

            language=None,

            page_count=None,

            character_count=len(
                text_content
            ),

            text_content=text_content,

            metadata={

                "author": metadata.author,

                "category": metadata.category,

                "comments": metadata.comments,

                "created": (
                    metadata.created.isoformat()
                    if metadata.created
                    else None
                ),

                "identifier": metadata.identifier,

                "keywords": metadata.keywords,

                "last_modified_by": metadata.last_modified_by,

                "last_printed": (
                    metadata.last_printed.isoformat()
                    if metadata.last_printed
                    else None
                ),

                "modified": (
                    metadata.modified.isoformat()
                    if metadata.modified
                    else None
                ),

                "revision": metadata.revision,

                "subject": metadata.subject,

                "version": metadata.version

            }

        )


docx_parser = DocxParser()