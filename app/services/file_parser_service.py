
import logging

import pandas as pd

from docx import (
    Document as DocxDocument
)

from pptx import Presentation

from pypdf import PdfReader

from app.services.ocr_service import (
    extract_text_from_image
)


logger = logging.getLogger(__name__)


# -----------------------------
# PDF
# -----------------------------

def parse_pdf(file_path: str):

    try:

        reader = PdfReader(file_path)

        pages = []

        for page_num, page in enumerate(
            reader.pages
        ):

            try:

                text = (
                    page.extract_text()
                    or ""
                )

            except Exception:

                logger.warning(
                    f"Failed to extract "
                    f"PDF page: "
                    f"{page_num + 1}"
                )

                text = ""

            if text.strip():

                pages.append({

                    "page_number": (
                        page_num + 1
                    ),

                    "text": text
                })

        return pages

    except Exception:

        logger.exception(
            f"PDF parse failed: "
            f"{file_path}"
        )

        raise


# -----------------------------
# DOCX
# -----------------------------

def parse_docx(file_path: str):

    try:

        doc = DocxDocument(
            file_path
        )

        text = "\n".join([

            p.text

            for p in doc.paragraphs

            if p.text.strip()
        ])

        return [{
            "page_number": 1,
            "text": text
        }]

    except Exception:

        logger.exception(
            f"DOCX parse failed: "
            f"{file_path}"
        )

        raise


# -----------------------------
# TXT
# -----------------------------

def parse_txt(file_path: str):

    try:

        with open(
            file_path,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as f:

            text = f.read()

        return [{
            "page_number": 1,
            "text": text
        }]

    except Exception:

        logger.exception(
            f"TXT parse failed: "
            f"{file_path}"
        )

        raise


# -----------------------------
# CSV
# -----------------------------

def parse_csv(file_path: str):

    try:

        df = pd.read_csv(
            file_path,
            nrows=5000
        )

        text = df.to_string()

        return [{
            "page_number": 1,
            "text": text
        }]

    except Exception:

        logger.exception(
            f"CSV parse failed: "
            f"{file_path}"
        )

        raise


# -----------------------------
# XLSX
# -----------------------------

def parse_xlsx(file_path: str):

    try:

        excel = pd.ExcelFile(
            file_path
        )

        all_text = []

        MAX_SHEETS = 20

        for sheet in (
            excel.sheet_names[
                :MAX_SHEETS
            ]
        ):

            try:

                df = pd.read_excel(
                    file_path,
                    sheet_name=sheet,
                    nrows=5000
                )

                all_text.append(
                    df.to_string()
                )

            except Exception:

                logger.warning(
                    f"Failed XLSX sheet: "
                    f"{sheet}"
                )

        return [{
            "page_number": 1,
            "text": "\n\n".join(
                all_text
            )
        }]

    except Exception:

        logger.exception(
            f"XLSX parse failed: "
            f"{file_path}"
        )

        raise


# -----------------------------
# PPTX
# -----------------------------

def parse_pptx(file_path: str):

    try:

        prs = Presentation(
            file_path
        )

        slides = []

        for idx, slide in enumerate(
            prs.slides
        ):

            slide_text = []

            for shape in slide.shapes:

                try:

                    if hasattr(
                        shape,
                        "text"
                    ):

                        text = (
                            shape.text
                            or ""
                        )

                        if text.strip():

                            slide_text.append(
                                text
                            )

                except Exception:

                    continue

            slides.append({

                "page_number": (
                    idx + 1
                ),

                "text": "\n".join(
                    slide_text
                )
            })

        return slides

    except Exception:

        logger.exception(
            f"PPTX parse failed: "
            f"{file_path}"
        )

        raise


# -----------------------------
# IMAGE OCR
# -----------------------------

async def parse_image(
    file_path: str
):

    try:

        text = (
            await extract_text_from_image(
                file_path
            )
        )

        return [{
            "page_number": 1,
            "text": text
        }]

    except Exception:

        logger.exception(
            f"Image parse failed: "
            f"{file_path}"
        )

        raise
